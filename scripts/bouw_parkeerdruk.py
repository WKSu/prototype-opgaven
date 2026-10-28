# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl", "pyproj", "pyshp"]
# ///
"""Bouwt data/parkeerdruk.js: gemeten parkeerdruk uit de parkeertellingen, per subbuurtdeel.

Leest de handmatig aangeleverde map data_manual/parkeerdruk/data: de zips van de tellingen
(Excel met resultaten per subbuurtdeel), de subbuurtdelen en de parkeersectoren (shapefiles in
RD New). Schrijft alleen aggregaten per subbuurtdeel, subbuurt en sector; geen kentekens.

    uv run scripts/bouw_parkeerdruk.py

Maatgevend moment: de werkdagnacht (meting op ma-vr, start tussen 23 en 5 uur). Heeft een
subbuurtdeel meerdere werkdagnachten, dan telt het gemiddelde; heeft het er geen, dan blijft de
druk leeg en staat het hoogste gemeten moment er alleen ter informatie bij.
"""

import argparse
import io
import json
import re
import zipfile
from itertools import pairwise
from pathlib import Path

import openpyxl
import shapefile
from pyproj import Transformer

ROOT = Path(__file__).resolve().parent.parent
GEBIEDEN = {
    4: "WK059904",
    6: "WK059906",
    14: "WK059914",
}  # zelfde codes als bouw_datalade
WERKDAGEN = {"maandag", "dinsdag", "woensdag", "donderdag", "vrijdag"}
VEREENVOUDIG_M = 8  # Douglas-Peucker in RD-meters; de kaart is zo’n 19 m per pixel
NODIG = [
    "Meting",
    "Dag",
    "Tijd",
    "Subbuurtdeel",
    "Buurtnaam",
    "Capaciteit Actueel Algemeen",
    "Bezetting Correct Algemeen",
    "Bezetting Fout",
]

RD = Transformer.from_crs(28992, 4326, always_xy=True)


def eis(obj, *sleutels: str, waar: str):
    # Een hernoemd veld moet luid falen, niet als lege kaart verschijnen.
    ontbreekt = [s for s in sleutels if s not in obj]
    if ontbreekt:
        raise KeyError(f"{waar}: velden ontbreken: {', '.join(ontbreekt)}")


def soort_moment(dag: str, tijd: str) -> str:
    dagen = {d.strip().lower() for d in dag.split("/")}
    start = int(re.match(r"\s*(\d{1,2})", tijd).group(1))
    deel = "nacht" if start >= 23 or start < 5 else "avond" if start >= 18 else "dag"
    return ("werkdag" if dagen <= WERKDAGEN else "weekend") + deel


def sbd_sleutel(waarde) -> tuple[int, int, int]:
    # '51.00', 68.1 en 53 betekenen buurt.subbuurt+deel: 51.00, 68.10, 53.00.
    b, rest = f"{float(waarde):.2f}".split(".")
    return int(b), int(rest[0]), int(rest[1])


def lees_telling(zip_pad: Path) -> dict:
    m = re.match(r"Parkeertelling (.+) \((\w+)\)$", zip_pad.stem)
    if not m:
        raise ValueError(f"{zip_pad.name}: naam niet als 'Parkeertelling X (code)'")
    with zipfile.ZipFile(zip_pad) as zf:
        xlsx = [n for n in zf.namelist() if n.endswith(".xlsx")]
        if len(xlsx) != 1:
            raise ValueError(f"{zip_pad.name}: verwacht één .xlsx, vond {xlsx}")
        wb = openpyxl.load_workbook(
            io.BytesIO(zf.read(xlsx[0])), read_only=True, data_only=True
        )
    eis(wb.sheetnames, "Terminologie", "Subbuurtdeelniveau", waar=xlsx[0])

    # De datum per meting staat alleen in Terminologie; dag en tijd nemen we uit de data.
    datum = {}
    for rij in wb["Terminologie"].iter_rows(values_only=True):
        cellen = [c for c in rij if c is not None]
        if len(cellen) >= 3 and re.fullmatch(r"M\d+", str(cellen[0])):
            d = cellen[2]
            datum[cellen[0]] = d.strftime("%d-%m-%Y") if hasattr(d, "strftime") else d

    rijen = wb["Subbuurtdeelniveau"].iter_rows(values_only=True)
    kop = next(rijen)
    eis(kop, *NODIG, waar=f"{xlsx[0]} Subbuurtdeelniveau")
    data = [dict(zip(kop, r)) for r in rijen if r[0]]
    wb.close()

    metingen = {}
    for r in data:
        metingen.setdefault(
            r["Meting"],
            {
                "id": r["Meting"],
                "dag": r["Dag"],
                "datum": datum.get(r["Meting"]),
                "tijd": r["Tijd"],
                "soort": soort_moment(r["Dag"], r["Tijd"]),
            },
        )
    return {
        "naam": m.group(1),
        "code": m.group(2),
        "bureau": xlsx[0].split("_")[1]
        if xlsx[0][:2].isdigit()
        else xlsx[0].split("_")[0],
        "metingen": sorted(metingen.values(), key=lambda x: int(x["id"][1:])),
        "rijen": data,
    }


def vereenvoudig(punten: list, tol: float) -> list:
    if len(punten) < 5:
        return punten
    # Een ring is gesloten (eerste = laatste punt): splits eerst op het verste punt, anders heeft
    # de basislijn lengte nul en valt er niets te vereenvoudigen.
    x0, y0 = punten[0]
    ver = max(
        range(len(punten)),
        key=lambda i: (punten[i][0] - x0) ** 2 + (punten[i][1] - y0) ** 2,
    )
    stapel, houd = [(0, ver), (ver, len(punten) - 1)], {0, ver, len(punten) - 1}
    while stapel:
        a, b = stapel.pop()
        (x1, y1), (x2, y2) = punten[a], punten[b]
        dx, dy = x2 - x1, y2 - y1
        n = (dx * dx + dy * dy) ** 0.5 or 1e-9
        verst, i_max = 0.0, None
        for i in range(a + 1, b):
            x, y = punten[i]
            d = abs(dy * x - dx * y + x2 * y1 - y2 * x1) / n
            if d > verst:
                verst, i_max = d, i
        if i_max is not None and verst > tol:
            houd.add(i_max)
            stapel += [(a, i_max), (i_max, b)]
    uit = [punten[i] for i in sorted(houd)]
    return uit if len(uit) >= 4 else punten


def ringen(vorm) -> list:
    grenzen = list(vorm.parts) + [len(vorm.points)]
    return [vorm.points[a:b] for a, b in pairwise(grenzen)]


def oppervlak(ring) -> float:
    return (
        sum(x1 * y2 - x2 * y1 for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]))
        / 2
    )


def geometrie(vorm) -> dict:
    # Shapefile: buitenringen met de klok mee, gaten ertegenin. Naar GeoJSON in WGS84.
    polygonen = []
    for ring in ringen(vorm):
        ring = vereenvoudig(ring, VEREENVOUDIG_M)
        lonlat = [[round(c, 5) for c in RD.transform(x, y)] for x, y in ring]
        if oppervlak(ring) <= 0 or not polygonen:
            polygonen.append([lonlat])
        else:
            polygonen[-1].append(lonlat)
    if len(polygonen) == 1:
        return {"type": "Polygon", "coordinates": polygonen[0]}
    return {"type": "MultiPolygon", "coordinates": polygonen}


def middelpunt(vorm) -> tuple[float, float]:
    ring = max(ringen(vorm), key=lambda r: abs(oppervlak(r)))
    a = oppervlak(ring) or 1e-9
    cx = sum(
        (x1 + x2) * (x1 * y2 - x2 * y1)
        for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1])
    )
    cy = sum(
        (y1 + y2) * (x1 * y2 - x2 * y1)
        for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1])
    )
    return cx / (6 * a), cy / (6 * a)


def in_ring(x: float, y: float, ring) -> bool:
    binnen = False
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]):
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            binnen = not binnen
    return binnen


def in_vorm(x: float, y: float, vorm) -> bool:
    # Even-oneven over alle ringen: een gat telt als buiten.
    return sum(in_ring(x, y, r) for r in ringen(vorm)) % 2 == 1


def druk(bez: float, cap: float):
    return round(100 * bez / cap, 1) if cap else None


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--map", type=Path, default=ROOT / "data_manual/parkeerdruk/data")
    ap.add_argument("--uit", type=Path, default=ROOT / "data/parkeerdruk.js")
    a = ap.parse_args()

    sbd_shp = shapefile.Reader(str(a.map / "Subbuurtdelen/Subbuurtdelen_vlakken.shp"))
    eis(
        [f[0] for f in sbd_shp.fields[1:]],
        "GEBIED",
        "BUURT",
        "SUBBUURT",
        "SBTDEEL",
        "TEKST",
        "BUURTNAAM",
        waar="Subbuurtdelen_vlakken",
    )
    sbd = {}
    for rec, vorm in zip(sbd_shp.records(), sbd_shp.shapes()):
        if rec["GEBIED"] in GEBIEDEN:
            sbd[(rec["BUURT"], rec["SUBBUURT"], rec["SBTDEEL"])] = (rec, vorm)

    # Per subbuurtdeel alle metingen uit alle tellingen; buiten het rayon valt weg.
    tellingen, gemeten = [], {}
    for z in sorted(a.map.glob("*.zip")):
        t = lees_telling(z)
        soort = {m["id"]: m["soort"] for m in t["metingen"]}
        binnen, som = set(), {}
        for r in t["rijen"]:
            k = sbd_sleutel(r["Subbuurtdeel"])
            if k not in sbd:
                continue
            if sbd[k][0]["BUURTNAAM"] != r["Buurtnaam"]:
                raise ValueError(
                    f"{t['code']}: {r['Subbuurtdeel']} is {r['Buurtnaam']!r}, shapefile zegt {sbd[k][0]['BUURTNAAM']!r}"
                )
            binnen.add(k)
            m = {
                "telling": len(tellingen),
                "meting": r["Meting"],
                "soort": soort[r["Meting"]],
                "cap": r["Capaciteit Actueel Algemeen"] or 0,
                "bez": (r["Bezetting Correct Algemeen"] or 0)
                + (r["Bezetting Fout"] or 0),
            }
            gemeten.setdefault(k, []).append(m)
            s = som.setdefault(r["Meting"], [0, 0])
            s[0] += m["bez"]
            s[1] += m["cap"]
        tellingen.append(
            {
                "naam": t["naam"],
                "code": t["code"],
                "bureau": t["bureau"],
                # Druk per meting over het hele getelde deel van het rayon: laat zien welk
                # moment het drukst is, en dus of de werkdagnacht maatgevend is.
                "metingen": [
                    {**m, "druk": druk(*som[m["id"]]) if m["id"] in som else None}
                    for m in t["metingen"]
                ],
                "subbuurtdelen": len(binnen),
                "gebieden": sorted({GEBIEDEN[sbd[k][0]["GEBIED"]] for k in binnen}),
            }
        )

    delen = []
    volgorde = sorted(sbd, key=lambda k: sbd[k][0]["TEKST"])
    for k in volgorde:
        rec, vorm = sbd[k]
        ms = gemeten.get(k, [])
        # Twee tellingen in hetzelfde subbuurtdeel komen in het rayon nu niet voor; gebeurt het
        # toch, dan middelen we over alle werkdagnachten van beide.
        nacht = [m for m in ms if m["soort"] == "werkdagnacht"]
        cap = sum(m["cap"] for m in nacht) / len(nacht) if nacht else None
        bez = sum(m["bez"] for m in nacht) / len(nacht) if nacht else None
        hoogste = max(
            (m for m in ms if m["cap"]), key=lambda m: m["bez"] / m["cap"], default=None
        )
        delen.append(
            {
                "code": rec["TEKST"],
                "subbuurt": rec["TEKST"][:5],
                "buurt": rec["BUURTNAAM"],
                "gebied": GEBIEDEN[rec["GEBIED"]],
                "telling": ms[0]["telling"] if ms else None,
                "nachten": len(nacht),
                "capaciteit": round(cap, 1) if cap is not None else None,
                "bezetting": round(bez, 1) if bez is not None else None,
                "druk": druk(bez, cap) if nacht else None,
                "perMeting": {m["meting"]: druk(m["bez"], m["cap"]) for m in ms},
                "hoogste": {
                    "meting": hoogste["meting"],
                    "soort": hoogste["soort"],
                    "druk": druk(hoogste["bez"], hoogste["cap"]),
                }
                if hoogste
                else None,
                "geometrie": geometrie(vorm),
            }
        )

    # Subbuurt = som van zijn delen. Alleen volledig als elk deel met capaciteit een nacht heeft.
    subbuurten = {}
    for d in delen:
        s = subbuurten.setdefault(
            d["subbuurt"],
            {
                "code": d["subbuurt"],
                "buurt": d["buurt"],
                "gebied": d["gebied"],
                "delen": 0,
                "geteld": 0,
                "cap": 0.0,
                "bez": 0.0,
                "zonderCap": 0,
            },
        )
        s["delen"] += 1
        if d["druk"] is not None:
            s["geteld"] += 1
            s["cap"] += d["capaciteit"]
            s["bez"] += d["bezetting"]
        elif d["nachten"] and not d["capaciteit"]:
            s["zonderCap"] += 1  # wel geteld, geen plekken: telt niet als ontbrekend
    sb_uit = []
    for s in subbuurten.values():
        sb_uit.append(
            {
                "code": s["code"],
                "buurt": s["buurt"],
                "gebied": s["gebied"],
                "delen": s["delen"],
                "geteld": s["geteld"],
                "volledig": s["geteld"] > 0
                and s["geteld"] + s["zonderCap"] == s["delen"],
                "capaciteit": round(s["cap"]) if s["geteld"] else None,
                "druk": druk(s["bez"], s["cap"]) if s["geteld"] else None,
            }
        )

    # Sectoren: alleen die een rayon-subbuurtdeel bevatten (op middelpunt). De gemeten nachtdruk
    # per sector is de som over de getelde delen waarvan het middelpunt in de sector ligt.
    sec_shp = shapefile.Reader(str(a.map / "parkeersectoren/ParkeerSectoren.shp"))
    eis([f[0] for f in sec_shp.fields[1:]], "SECTOR_ID", waar="ParkeerSectoren")
    punten = [(d, middelpunt(sbd[k][1])) for k, d in zip(volgorde, delen)]
    sectoren = []
    for rec, vorm in zip(sec_shp.records(), sec_shp.shapes()):
        x0, y0, x1, y1 = vorm.bbox
        erin = [
            d
            for d, (x, y) in punten
            if x0 <= x <= x1 and y0 <= y <= y1 and in_vorm(x, y, vorm)
        ]
        if not erin:
            continue
        geteld = [d for d in erin if d["druk"] is not None]
        cap = sum(d["capaciteit"] for d in geteld)
        sectoren.append(
            {
                "id": str(rec["SECTOR_ID"]),
                "subbuurtdelen": len(erin),
                "delen": [d["code"] for d in erin],
                "geteld": len(geteld),
                "nachtdruk": druk(sum(d["bezetting"] for d in geteld), cap)
                if geteld
                else None,
                "geometrie": geometrie(vorm),
            }
        )

    uit = {
        "bron": "Parkeertellingen gemeente Rotterdam (handmatig aangeleverd), subbuurtdelen en parkeersectoren (shapefiles)",
        "regel": "bruto parkeerdruk = (correct geparkeerd op algemene plekken + foutgeparkeerd) ÷ actuele capaciteit algemene plekken",
        "maatgevend": "werkdagnacht: meting op maandag t/m vrijdag die start tussen 23 en 5 uur; bij meer nachten het gemiddelde",
        "tellingen": tellingen,
        "subbuurtdelen": delen,
        "subbuurten": sorted(sb_uit, key=lambda s: s["code"]),
        "sectoren": sorted(sectoren, key=lambda s: s["id"]),
    }
    tekst = json.dumps(uit, ensure_ascii=False, separators=(",", ":"))
    a.uit.write_text(
        "// Gegenereerd door scripts/bouw_parkeerdruk.py. Niet met de hand bewerken.\n"
        f"const PARKEERDRUK = {tekst};\n",
        encoding="utf-8",
    )
    print(f"{a.uit}: {len(tekst) / 1024:.0f} kB")
    for t in tellingen:
        print(
            f"  {t['code']} {t['naam']}: {t['subbuurtdelen']} subbuurtdelen in het rayon"
        )


if __name__ == "__main__":
    main()
