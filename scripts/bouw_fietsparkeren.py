# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl", "pyproj", "pyshp"]
# ///
"""Bouwt data/fietsparkeren.js: capaciteit en fietsparkeerdruk uit de fietsparkeertellingen.

Leest de handmatig aangeleverde zips in data_manual/fietsparkeertellingen (Ecorys, 2022-2023):
per telling een Excel met resultaten per subbuurt en een puntenbestand van de voorzieningen.
De subbuurtcodes uit de telling (buurt_subbuurt, '63_1') worden via de subbuurtdelen
(data_manual/parkeerdruk/data/Subbuurtdelen) vertaald naar de vijfcijferige code van de viewer.

    uv run scripts/bouw_fietsparkeren.py

Fietsparkeerdruk = alle gestalde fietsen (in, nabij en buiten een voorziening) ÷ capaciteit van
de voorzieningen. Maatgevend is per telling het drukste van de vier meetmomenten.
"""

import argparse
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import shapefile
from bouw_parkeerdruk import GEBIEDEN, RD, ROOT, eis

KOLOM = {
    "in": "Aantal geparkeerde fietsen in voorziening",
    "nabij": "Aantal geparkeerde fietsen nabij voorziening",
    "wild": "Aantal wildgeparkeerde fietsen",
}
NODIG = [
    "Subbuurt",
    "Telmoment",
    "Datum",
    "Tijdsperiode",
    "Capaciteit voorzieningen",
    *KOLOM.values(),
]
# Moment → voorvoegsel van de velden in het puntenbestand (DM_in, DA_aantal, ...).
VOORVOEGSEL = {
    "Dinsdagmiddag": "DM",
    "Dinsdagavond": "DA",
    "Zaterdagmiddag": "ZM",
    "Zaterdagavond": "ZA",
}


def pct(teller: float, noemer: float):
    return round(100 * teller / noemer, 1) if noemer else None


def lees_telling(zip_pad: Path, codes: dict) -> dict:
    m = re.match(r"Fietsparkeertelling (.+) \((\w+)\)$", zip_pad.stem)
    if not m:
        raise ValueError(
            f"{zip_pad.name}: naam niet als 'Fietsparkeertelling X (code)'"
        )
    with zipfile.ZipFile(zip_pad) as zf:
        namen = [n for n in zf.namelist() if not n.endswith(":Zone.Identifier")]
        xlsx = [n for n in namen if n.endswith(".xlsx")]
        if len(xlsx) != 1:
            raise ValueError(f"{zip_pad.name}: verwacht één .xlsx, vond {xlsx}")
        wb = openpyxl.load_workbook(
            io.BytesIO(zf.read(xlsx[0])), read_only=True, data_only=True
        )
        eis(wb.sheetnames, "Subbuurtniveau", waar=xlsx[0])
        rijen = wb["Subbuurtniveau"].iter_rows(values_only=True)
        kop = next(rijen)
        eis(kop, *NODIG, waar=f"{xlsx[0]} Subbuurtniveau")
        data = [dict(zip(kop, r)) for r in rijen if r[0]]
        wb.close()
        voorz = [n for n in namen if n.endswith("_Voorzieningen.shp")]
        if len(voorz) != 1:
            raise ValueError(
                f"{zip_pad.name}: verwacht één Voorzieningen.shp, vond {voorz}"
            )
        basis = voorz[0][:-4]
        punten = shapefile.Reader(
            shp=io.BytesIO(zf.read(basis + ".shp")),
            shx=io.BytesIO(zf.read(basis + ".shx")),
            dbf=io.BytesIO(zf.read(basis + ".dbf")),
        )
        velden = [f[0] for f in punten.fields[1:]]
        eis(velden, "Capaciteit", "Type", "Sectie_ID", waar=basis)
        voorzieningen = [
            (r.as_dict(), s.points[0])
            for r, s in zip(punten.records(), punten.shapes())
            if s.points
        ]

    metingen = {}
    for r in data:
        metingen.setdefault(
            r["Telmoment"],
            {
                "moment": r["Telmoment"],
                "datum": r["Datum"].strftime("%d-%m-%Y")
                if hasattr(r["Datum"], "strftime")
                else str(r["Datum"]),
                "tijd": r["Tijdsperiode"],
                "cap": 0,
                "fietsen": 0,
                "in": 0,
            },
        )
        mt = metingen[r["Telmoment"]]
        mt["cap"] += r["Capaciteit voorzieningen"] or 0
        mt["fietsen"] += sum(r[k] or 0 for k in KOLOM.values())
        mt["in"] += r[KOLOM["in"]] or 0
    onbekend = set(metingen) - set(VOORVOEGSEL)
    if onbekend:
        raise ValueError(f"{zip_pad.name}: onbekende telmomenten {sorted(onbekend)}")
    drukst = max(
        metingen.values(), key=lambda x: x["fietsen"] / x["cap"] if x["cap"] else 0
    )

    subbuurten = {}
    for r in data:
        code = codes.get(r["Subbuurt"])
        if code is None:
            raise ValueError(
                f"{zip_pad.name}: subbuurt {r['Subbuurt']} niet in de subbuurtdelen"
            )
        s = subbuurten.setdefault(
            code,
            {
                "code": code,
                "capaciteit": r["Capaciteit voorzieningen"] or 0,
                "perMoment": {},
            },
        )
        n = {k: r[c] or 0 for k, c in KOLOM.items()}
        f = sum(n.values())
        s["perMoment"][r["Telmoment"]] = {
            **n,
            "fietsen": f,
            "druk": pct(f, s["capaciteit"]),
            "bezetting": pct(n["in"], s["capaciteit"]),
        }

    vv = VOORVOEGSEL[drukst["moment"]]
    punten_uit = []
    for rec, (x, y) in voorzieningen:
        lon, lat = RD.transform(x, y)
        cap = int(float(rec["Capaciteit"] or 0))
        aantal = rec.get(f"{vv}_aantal")
        punten_uit.append(
            {
                "lon": round(lon, 5),
                "lat": round(lat, 5),
                "capaciteit": cap,
                "type": rec["Type"],
                "fietsen": round(float(aantal)) if aantal not in (None, "") else None,
            }
        )
    return {
        "naam": m.group(1),
        "code": m.group(2),
        "maatgevend": drukst["moment"],
        "metingen": [
            {
                "moment": x["moment"],
                "datum": x["datum"],
                "tijd": x["tijd"],
                "capaciteit": x["cap"],
                "fietsen": x["fietsen"],
                "druk": pct(x["fietsen"], x["cap"]),
                "bezetting": pct(x["in"], x["cap"]),
            }
            for x in sorted(
                metingen.values(), key=lambda x: list(VOORVOEGSEL).index(x["moment"])
            )
        ],
        "subbuurten": subbuurten,
        "voorzieningen": punten_uit,
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument(
        "--map", type=Path, default=ROOT / "data_manual/fietsparkeertellingen"
    )
    ap.add_argument(
        "--sbd",
        type=Path,
        default=ROOT
        / "data_manual/parkeerdruk/data/Subbuurtdelen/Subbuurtdelen_vlakken.shp",
    )
    ap.add_argument("--uit", type=Path, default=ROOT / "data/fietsparkeren.js")
    a = ap.parse_args()

    # '63_1' → buurt 63, subbuurt 1 → gebied + buurt + subbuurt = '14631'. Buurtnummers zijn uniek.
    sbd = shapefile.Reader(str(a.sbd))
    eis(
        [f[0] for f in sbd.fields[1:]],
        "GEBIED",
        "BUURT",
        "SUBBUURT",
        "TEKST",
        waar="Subbuurtdelen_vlakken",
    )
    codes, gebied_van = {}, {}
    for r in sbd.records():
        codes[f"{r['BUURT']}_{r['SUBBUURT']}"] = r["TEKST"][:5]
        gebied_van[r["TEKST"][:5]] = r["GEBIED"]

    tellingen, subbuurten, voorzieningen = [], [], []
    for z in sorted(a.map.glob("*.zip")):
        t = lees_telling(z, codes)
        i = len(tellingen)
        buiten = [c for c in t["subbuurten"] if gebied_van[c] not in GEBIEDEN]
        if buiten:
            raise ValueError(f"{z.name}: subbuurten buiten het rayon: {buiten}")
        for s in t["subbuurten"].values():
            p = s["perMoment"][t["maatgevend"]]
            subbuurten.append(
                {
                    "code": s["code"],
                    "gebied": GEBIEDEN[gebied_van[s["code"]]],
                    "telling": i,
                    "capaciteit": s["capaciteit"],
                    "fietsen": p["fietsen"],
                    "wild": p["wild"],
                    "druk": p["druk"],
                    "bezetting": p["bezetting"],
                    "tekort": max(0, p["fietsen"] - s["capaciteit"]),
                    "perMoment": {k: v["druk"] for k, v in s["perMoment"].items()},
                }
            )
        voorzieningen += [{**v, "telling": i} for v in t["voorzieningen"]]
        tellingen.append(
            {
                "naam": t["naam"],
                "code": t["code"],
                "jaar": t["metingen"][0]["datum"][-4:],
                "gebied": GEBIEDEN[gebied_van[next(iter(t["subbuurten"]))]],
                "subbuurten": len(t["subbuurten"]),
                "maatgevend": t["maatgevend"],
                "metingen": t["metingen"],
            }
        )

    dubbel = {
        s["code"]
        for s in subbuurten
        if sum(x["code"] == s["code"] for x in subbuurten) > 1
    }
    if dubbel:
        raise ValueError(f"subbuurten in meer dan één telling: {sorted(dubbel)}")

    uit = {
        "bron": "Fietsparkeertellingen gemeente Rotterdam, uitgevoerd door Ecorys, 2022–2023 (handmatig aangeleverd)",
        "regel": "fietsparkeerdruk = alle gestalde fietsen (in, nabij en buiten een voorziening) ÷ capaciteit van de voorzieningen; bezetting = fietsen in een voorziening ÷ capaciteit",
        "maatgevend": "per telling het drukste van de vier meetmomenten (dinsdag en zaterdag, middag en avond)",
        "tellingen": tellingen,
        "subbuurten": sorted(subbuurten, key=lambda s: s["code"]),
        "voorzieningen": voorzieningen,
    }
    tekst = json.dumps(uit, ensure_ascii=False, separators=(",", ":"))
    a.uit.write_text(
        "// Gegenereerd door scripts/bouw_fietsparkeren.py. Niet met de hand bewerken.\n"
        f"const FIETSPARKEREN = {tekst};\n",
        encoding="utf-8",
    )
    print(
        f"{a.uit}: {len(tekst) / 1024:.0f} kB, {len(subbuurten)} subbuurten, {len(voorzieningen)} voorzieningen"
    )
    for t in tellingen:
        m = next(x for x in t["metingen"] if x["moment"] == t["maatgevend"])
        print(
            f"  {t['code']} {t['naam']}: {t['maatgevend']}, druk {m['druk']}%, capaciteit {m['capaciteit']}"
        )


if __name__ == "__main__":
    main()
